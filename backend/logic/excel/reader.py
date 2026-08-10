from openpyxl import load_workbook


class ExcelReader:

    def __init__(self, file):

        self.file = file

        self.workbook = load_workbook(
            filename=file,
            data_only=True
        )

    # -------------------------------------------------
    # Workbook
    # -------------------------------------------------

    def get_workbook(self):

        return self.workbook

    # -------------------------------------------------
    # Sheet names
    # -------------------------------------------------

    def get_sheet_names(self):

        return self.workbook.sheetnames

    # -------------------------------------------------
    # Worksheet
    # -------------------------------------------------

    def get_sheet(self, name=None):

        if name is None:
            return self.workbook.active

        return self.workbook[name]

    # -------------------------------------------------
    # Read values
    # -------------------------------------------------

    def read_sheet(self, name=None):

        ws = self.get_sheet(name)

        rows = []

        for row in ws.iter_rows(values_only=True):

            rows.append(list(row))

        return rows

    # -------------------------------------------------
    # Merged cells
    # -------------------------------------------------

    def get_merged_cells(self, name=None):

        ws = self.get_sheet(name)

        return [
            str(rng)
            for rng in ws.merged_cells.ranges
        ]

    # -------------------------------------------------
    # Dimensions
    # -------------------------------------------------

    def get_size(self, name=None):

        ws = self.get_sheet(name)

        return {

            "rows": ws.max_row,
            "columns": ws.max_column

        }

    # -------------------------------------------------
    # Metadata
    # -------------------------------------------------

    def info(self):

        worksheets = []

        for sheet_name in self.get_sheet_names():

            ws = self.get_sheet(sheet_name)

            worksheets.append({

                "name": sheet_name,
                "rows": ws.max_row,
                "columns": ws.max_column,
                "merged_cells": len(ws.merged_cells.ranges)

            })

        return {

            "sheet_names": self.get_sheet_names(),
            "worksheets": worksheets

        }