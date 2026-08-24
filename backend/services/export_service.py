# -*- coding: utf-8 -*-
"""
Service d'export centralisé pour le Centre Import/Export FTTH.
Support des formats Excel, CSV, PDF.
Permet la sélection des colonnes, filtres, modèles.
"""
import csv
import io
import json
import logging
import os
import re
import tempfile
import time
from datetime import datetime, date, timedelta
from typing import Optional, List, Dict, Any
from zipfile import ZipFile

from sqlalchemy import select, and_, func, or_
from sqlalchemy.ext.asyncio import AsyncSession
from openpyxl import Workbook
from openpyxl.styles import PatternFill, Font, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from backend.database.models import (
    Job, JobStatus, JobType,
    Technician, Assignment, Orienteur, Sector,
    ExportTemplate, ExportHistory, User, JobActivityLog,
)
from backend.config import get_settings
from backend.logic.job_sectors import hydrate_job_sector_identities

logger = logging.getLogger(__name__)


_CIVIL_DATE_PATTERN = re.compile(r"^\d{4}-\d{2}-\d{2}$")


def _civil_filter_boundary(value, *, end: bool = False):
    """Normalize a civil date to an index-friendly datetime boundary.

    Timestamp inputs keep their exact historical inclusive semantics. A
    ``YYYY-MM-DD``/``date`` end bound becomes the exclusive start of the next
    day so an export never drops records after midnight on its final day.
    """

    civil_date = None
    if isinstance(value, datetime):
        return value, False
    if isinstance(value, date):
        civil_date = value
    elif isinstance(value, str) and _CIVIL_DATE_PATTERN.fullmatch(value.strip()):
        try:
            civil_date = date.fromisoformat(value.strip())
        except ValueError:
            civil_date = None

    if civil_date is None:
        return value, False

    if end:
        civil_date += timedelta(days=1)
    return datetime.combine(civil_date, datetime.min.time()), True


def _enum_filter_value(enum_type, value, *, label: str):
    if isinstance(value, enum_type):
        return value

    token = str(value if value is not None else "").strip()
    if not token:
        raise ValueError(f"Filtre {label} vide.")

    for candidate in (token, token.lower(), token.upper()):
        try:
            return enum_type(candidate)
        except ValueError:
            pass

    member = enum_type.__members__.get(token.upper())
    if member is not None:
        return member

    raise ValueError(f"Filtre {label} invalide : {token}.")

# ── CHARTE GRAPHIQUE ─────────────────────────────────────────────
BLEU_NAVY = "1F497D"
BLANC = "FFFFFF"
GRIS_CLAIR = "F2F2F2"
GRIS_BORDURE = "D9D9D9"
VERT_VIF = "00B050"
ORANGE = "FFC000"
ROUGE = "C00000"

# ── COLONNES DISPONIBLES POUR L'EXPORT ──────────────────────────
EXPORT_COLUMNS = {
    "secteur": {
        "label": "Secteur",
        "field": "secteur",
        "default": True,
        "category": "localisation",
    },
    "date": {
        "label": "Date",
        "field": "date",
        "default": True,
        "category": "dates",
    },
    "commande": {
        "label": "Commande",
        "field": "commande",
        "default": True,
        "category": "identification",
    },
    "client": {
        "label": "Client",
        "field": "client",
        "default": True,
        "category": "client",
    },
    "contact": {
        "label": "Contact",
        "field": "contact",
        "default": True,
        "category": "client",
    },
    "adresse": {
        "label": "Adresse",
        "field": "adresse",
        "default": True,
        "category": "localisation",
    },
    "operateur": {
        "label": "Opérateur",
        "field": "operateur",
        "default": True,
        "category": "identification",
    },
    "nro": {
        "label": "NRO",
        "field": "nro",
        "default": False,
        "category": "reseau",
    },
    "sro": {
        "label": "SRO",
        "field": "sro",
        "default": False,
        "category": "reseau",
    },
    "pbo": {
        "label": "PBO",
        "field": "pbo",
        "default": False,
        "category": "reseau",
    },
    "pto": {
        "label": "PTO",
        "field": "pto",
        "default": False,
        "category": "reseau",
    },
    "splitter": {
        "label": "Splitter",
        "field": "splitter",
        "default": False,
        "category": "reseau",
    },
    "port": {
        "label": "Port",
        "field": "port",
        "default": False,
        "category": "reseau",
    },
    "pco": {
        "label": "PCO",
        "field": "pco",
        "default": False,
        "category": "reseau",
    },
    "gps_client": {
        "label": "GPS Client",
        "field": "gps_client",
        "default": False,
        "category": "gps",
    },
    "gps_pbo": {
        "label": "GPS PBO",
        "field": "gps_pbo",
        "default": False,
        "category": "gps",
    },
    "gps_splitter": {
        "label": "GPS Splitter",
        "field": "gps_splitter",
        "default": False,
        "category": "gps",
    },
    "gps_derivation": {
        "label": "GPS Dérivation",
        "field": "gps_derivation",
        "default": False,
        "category": "gps",
    },
    "signal": {
        "label": "Signal",
        "field": "signal",
        "default": True,
        "category": "technique",
    },
    "statut": {
        "label": "Statut",
        "field": "statut",
        "default": True,
        "category": "suivi",
    },
    "technicien": {
        "label": "Technicien",
        "field": "technicien",
        "default": True,
        "category": "ressources",
    },
    "chef_orienteur": {
        "label": "Chef Orienteur",
        "field": "chef_orienteur",
        "default": False,
        "category": "ressources",
    },
    "heure_debut": {
        "label": "Heure début",
        "field": "heure_debut",
        "default": False,
        "category": "temps",
    },
    "heure_fin": {
        "label": "Heure fin",
        "field": "heure_fin",
        "default": False,
        "category": "temps",
    },
    "temps_intervention": {
        "label": "Temps intervention",
        "field": "temps_intervention",
        "default": False,
        "category": "temps",
    },
    "puissance_optique": {
        "label": "Puissance Optique",
        "field": "puissance_optique",
        "default": False,
        "category": "technique",
    },
    "longueur_cable": {
        "label": "Longueur câble",
        "field": "longueur_cable",
        "default": False,
        "category": "technique",
    },
    "sn_ont": {
        "label": "SN ONT",
        "field": "sn_ont",
        "default": False,
        "category": "equipement",
    },
    "sn_routeur": {
        "label": "SN Routeur",
        "field": "sn_routeur",
        "default": False,
        "category": "equipement",
    },
    "sn_wifi": {
        "label": "SN WiFi",
        "field": "sn_wifi",
        "default": False,
        "category": "equipement",
    },
    "mac": {
        "label": "MAC",
        "field": "mac",
        "default": False,
        "category": "equipement",
    },
    "photos": {
        "label": "Photos",
        "field": "photos",
        "default": False,
        "category": "multimedia",
    },
    "signature": {
        "label": "Signature",
        "field": "signature",
        "default": False,
        "category": "multimedia",
    },
    "commentaires": {
        "label": "Commentaires",
        "field": "commentaires",
        "default": False,
        "category": "notes",
    },
    "observation": {
        "label": "Observation",
        "field": "observation",
        "default": False,
        "category": "notes",
    },
    "remarque": {
        "label": "Remarque",
        "field": "remarque",
        "default": False,
        "category": "notes",
    },
    "type_intervention": {
        "label": "Type d'intervention",
        "field": "type_intervention",
        "default": True,
        "category": "identification",
    },
    "priorite": {
        "label": "Priorité",
        "field": "priorite",
        "default": False,
        "category": "suivi",
    },
    "validation_status": {
        "label": "Validation",
        "field": "validation_status",
        "default": False,
        "category": "suivi",
    },
}

# ── PROFILS D'EXPORT PRÉDÉFINIS ─────────────────────────────────
EXPORT_PROFILES = {
    "export_orange": {
        "name": "Export Orange",
        "columns": [
            "date", "commande", "client", "contact", "adresse",
            "nro", "sro", "pbo", "pto", "splitter", "port",
            "signal", "statut", "technicien",
            "sn_ont", "sn_routeur", "mac",
            "puissance_optique", "longueur_cable",
            "commentaires",
        ],
        "filters": {"operator": "ORANGE"},
    },
    "export_iam": {
        "name": "Export IAM",
        "columns": [
            "date", "commande", "client", "contact", "adresse",
            "nro", "sro", "pbo", "pto",
            "statut", "technicien",
            "sn_ont", "sn_routeur",
            "commentaires",
        ],
        "filters": {"operator": "IAM"},
    },
    "export_inwi": {
        "name": "Export Inwi",
        "columns": [
            "date", "commande", "client", "contact", "adresse",
            "splitter", "port",
            "statut", "technicien",
            "sn_ont", "mac",
            "commentaires",
        ],
        "filters": {"operator": "INWI"},
    },
    "export_facturation": {
        "name": "Export Facturation",
        "columns": [
            "date", "commande", "client", "operateur",
            "type_intervention", "statut", "technicien",
            "heure_debut", "heure_fin", "temps_intervention",
            "sn_ont", "sn_routeur",
            "commentaires",
        ],
        "filters": {},
    },
    "export_controle_qualite": {
        "name": "Export Contrôle Qualité",
        "columns": [
            "date", "commande", "client", "operateur",
            "type_intervention", "statut", "technicien",
            "signal", "puissance_optique",
            "photos", "signature",
            "validation_status", "commentaires",
        ],
        "filters": {"status": ["completed"]},
    },
    "export_chef_orienteur": {
        "name": "Export Chef Orienteur",
        "columns": [
            "secteur", "date", "commande", "client",
            "operateur", "type_intervention", "statut",
            "technicien", "chef_orienteur",
            "heure_debut", "heure_fin", "temps_intervention",
            "priorite", "commentaires",
        ],
        "filters": {},
    },
}


def _thin_border() -> Border:
    side = Side(style="thin", color=GRIS_BORDURE)
    return Border(left=side, right=side, top=side, bottom=side)


def _header_fill(hex_color: str) -> PatternFill:
    return PatternFill("solid", fgColor=hex_color)


def _header_font(bold: bool = True, color: str = BLANC) -> Font:
    return Font(bold=bold, color=color, name="Calibri", size=10)


def _center_align() -> Alignment:
    return Alignment(horizontal="center", vertical="center", wrap_text=True)


def _left_align() -> Alignment:
    return Alignment(horizontal="left", vertical="center", wrap_text=True)


def _autofit_columns(ws, min_width: int = 12, max_width: int = 45):
    """Ajuste automatiquement la largeur des colonnes."""
    for col in ws.columns:
        max_length = 0
        col_letter = get_column_letter(col[0].column)
        for cell in col:
            if cell.value:
                max_length = max(max_length, len(str(cell.value)))
        adjusted = min(max(max_length + 2, min_width), max_width)
        ws.column_dimensions[col_letter].width = adjusted


class FieldOptExportService:
    """Service centralisé d'export des données FieldOpt."""

    # ── BUILDERS DE JEUX DE DONNÉES ──────────────────────────

    @staticmethod
    async def _extract_job_value(job: Job, field: str, db: AsyncSession = None) -> Any:
        """Extrait la valeur d'un champ du modèle Job pour l'export."""
        mapping = {
            "secteur": lambda j: (
                getattr(j, "_canonical_sector_name", None)
                or getattr(j.__dict__.get("sector"), "name", None)
                or ""
            ),
            "date": lambda j: j.scheduled_date.strftime("%d/%m/%Y") if j.scheduled_date else "",
            "commande": lambda j: j.job_number or "",
            "client": lambda j: j.customer_name or "",
            "contact": lambda j: j.customer_phone or "",
            "adresse": lambda j: j.service_address or "",
            "operateur": lambda j: j.operator or "",
            "nro": lambda j: j.nro_raw or "",
            "sro": lambda j: j.sro_raw or "",
            "pbo": lambda j: j.pbo_raw or "",
            "pto": lambda j: j.pto_raw or "",
            "splitter": lambda j: j.splitter_raw or "",
            "port": lambda j: str(j.splitter_port_raw) if j.splitter_port_raw else "",
            "pco": lambda j: "",
            "gps_client": lambda j: (
                f"{j.latitude:.6f}, {j.longitude:.6f}"
                if j.latitude is not None and j.longitude is not None
                else ""
            ),
            "gps_pbo": lambda j: "",
            "gps_splitter": lambda j: "",
            "gps_derivation": lambda j: "",
            "signal": lambda j: f"{j.optical_power_dbm} dBm" if j.optical_power_dbm else "",
            "statut": lambda j: j.status.value if j.status else "",
            "technicien": lambda j: j.assigned_technician_name or "",
            "chef_orienteur": lambda j: j.orienteur.name if j.orienteur else "",
            "heure_debut": lambda j: j.started_at.strftime("%H:%M") if j.started_at else "",
            "heure_fin": lambda j: j.completed_at.strftime("%H:%M") if j.completed_at else "",
            "temps_intervention": lambda j: f"{j.real_duration_minutes} min" if j.real_duration_minutes else "",
            "puissance_optique": lambda j: f"{j.optical_power_dbm} dBm" if j.optical_power_dbm else "",
            "longueur_cable": lambda j: f"{j.cable_length_m} m" if j.cable_length_m else "",
            "sn_ont": lambda j: j.ont_serial or "",
            "sn_routeur": lambda j: j.router_serial or "",
            "sn_wifi": lambda j: j.wifi_box_serial or "",
            "mac": lambda j: j.mac_address or "",
            "photos": lambda j: "Oui" if (j.before_photo or j.after_photo) else "Non",
            "signature": lambda j: "Oui" if j.client_signature else "Non",
            "commentaires": lambda j: j.notes or "",
            "observation": lambda j: j.coordinator_comments or "",
            "remarque": lambda j: j.description or "",
            "type_intervention": lambda j: j.job_type.value if j.job_type else "",
            "priorite": lambda j: j.priority.value if j.priority else "",
            "validation_status": lambda j: j.validation_status or "",
        }
        extractor = mapping.get(field)
        if extractor:
            try:
                return extractor(job)
            except Exception:
                return ""
        return ""

    @staticmethod
    async def build_job_query(
        db: AsyncSession,
        filters: dict = None,
    ):
        """
        Construit la requête SQLAlchemy pour récupérer les jobs selon les filtres.
        Optimisée pour ne charger que les données nécessaires.
        """
        query = (
            select(Job)
            .outerjoin(
                Assignment,
                and_(
                    Assignment.job_id == Job.id,
                    Assignment.ended_at.is_(None),
                ),
            )
            .outerjoin(Technician)
            .outerjoin(Orienteur)
            .where(Job.deleted_at.is_(None))
        )

        if not filters:
            filters = {}

        # Filtre par période
        start_date = filters.get("start_date")
        end_date = filters.get("end_date")
        if start_date:
            normalized_start, _ = _civil_filter_boundary(start_date)
            query = query.where(Job.scheduled_date >= normalized_start)
        if end_date:
            normalized_end, end_is_civil = _civil_filter_boundary(
                end_date,
                end=True,
            )
            query = query.where(
                Job.scheduled_date < normalized_end
                if end_is_civil
                else Job.scheduled_date <= normalized_end
            )

        # Filtre par date relative
        date_preset = filters.get("date_preset")
        if date_preset:
            today = datetime.utcnow().date()
            if date_preset == "today":
                query = query.where(func.date(Job.scheduled_date) == today)
            elif date_preset == "yesterday":
                yesterday = today - timedelta(days=1)
                query = query.where(func.date(Job.scheduled_date) == yesterday)
            elif date_preset == "this_week":
                start_week = today - timedelta(days=today.weekday())
                query = query.where(func.date(Job.scheduled_date) >= start_week)
            elif date_preset == "this_month":
                start_month = today.replace(day=1)
                query = query.where(func.date(Job.scheduled_date) >= start_month)
            elif date_preset == "last_month":
                first_prev = (today.replace(day=1) - timedelta(days=1)).replace(day=1)
                last_prev = today.replace(day=1) - timedelta(days=1)
                query = query.where(
                    and_(
                        func.date(Job.scheduled_date) >= first_prev,
                        func.date(Job.scheduled_date) <= last_prev,
                    )
                )

        # Filtre par opérateur
        operator = filters.get("operator")
        if operator and operator != "TOUS":
            normalized_operator = str(operator).strip().upper()
            query = query.where(func.upper(func.trim(Job.operator)) == normalized_operator)

        # Filtre par statut
        status_filter = filters.get("status")
        if status_filter:
            if isinstance(status_filter, list) and len(status_filter) > 0:
                query = query.where(
                    Job.status.in_([
                        _enum_filter_value(JobStatus, value, label="statut")
                        for value in status_filter
                    ])
                )
            elif str(status_filter).strip().upper() != "TOUTES":
                query = query.where(
                    Job.status
                    == _enum_filter_value(JobStatus, status_filter, label="statut")
                )

        # ``sector_id`` is applied after the shared canonical hydration in
        # ``get_filtered_jobs``.  Keeping that reconciliation out of ad-hoc
        # SQL guarantees that legacy routing codes and structured territories
        # follow exactly the same contract as Job API responses.

        # Filtre par technicien
        technician_id = filters.get("technician_id")
        if technician_id:
            query = query.where(Assignment.technician_id == technician_id)

        # Filtre par orienteur
        orienteur_id = filters.get("orienteur_id")
        if orienteur_id:
            query = query.where(Technician.orienteur_id == orienteur_id)

        # Filtre par type d'intervention
        job_type = filters.get("job_type")
        if job_type:
            query = query.where(
                Job.job_type
                == _enum_filter_value(JobType, job_type, label="type")
            )

        # Recherche textuelle
        search = filters.get("search")
        if search:
            search_term = f"%{search}%"
            query = query.where(
                or_(
                    Job.customer_name.ilike(search_term),
                    Job.job_number.ilike(search_term),
                    Job.customer_phone.ilike(search_term),
                    Job.service_address.ilike(search_term),
                    Job.ont_serial.ilike(search_term),
                    Job.pto_raw.ilike(search_term),
                    Job.pbo_raw.ilike(search_term),
                    Job.nro_raw.ilike(search_term),
                    Job.assigned_technician_name.ilike(search_term),
                    Job.notes.ilike(search_term),
                )
            )

        # Tri par date
        query = query.order_by(Job.scheduled_date.desc(), Job.id.desc())

        return query

    @staticmethod
    async def get_filtered_jobs(
        db: AsyncSession,
        filters: dict = None,
        *,
        limit: int | None = None,
    ) -> List[Job]:
        """Return export jobs after applying the canonical sector contract."""

        query_filters = dict(filters or {})
        requested_sector_id = query_filters.pop("sector_id", None)
        try:
            requested_sector_id = (
                int(requested_sector_id)
                if requested_sector_id not in (None, "")
                else None
            )
        except (TypeError, ValueError) as exc:
            raise ValueError(
                f"Filtre secteur invalide : {requested_sector_id}."
            ) from exc
        if requested_sector_id is not None and requested_sector_id <= 0:
            raise ValueError(f"Filtre secteur invalide : {requested_sector_id}.")

        query = await FieldOptExportService.build_job_query(db, query_filters)
        result = await db.execute(query)
        jobs = list(result.scalars().all())
        await hydrate_job_sector_identities(db, jobs)

        if requested_sector_id is not None:
            jobs = [
                job
                for job in jobs
                if getattr(job, "_canonical_sector_id", None)
                == requested_sector_id
            ]

        return jobs[:limit] if limit is not None else jobs

    @staticmethod
    async def build_job_dict(
        job: Job,
        columns: List[str],
        db: AsyncSession = None,
    ) -> dict:
        """Construit un dictionnaire de valeurs pour un job selon les colonnes sélectionnées."""
        row = {}
        for col_key in columns:
            label = EXPORT_COLUMNS.get(col_key, {}).get("label", col_key)
            value = await FieldOptExportService._extract_job_value(job, col_key, db)
            row[label] = value
        return row

    @staticmethod
    async def get_jobs_data(
        db: AsyncSession,
        filters: dict = None,
        columns: List[str] = None,
    ) -> List[dict]:
        """Récupère les données des jobs sous forme de dictionnaires."""
        if columns is None:
            columns = [k for k, v in EXPORT_COLUMNS.items() if v["default"]]

        jobs = await FieldOptExportService.get_filtered_jobs(db, filters)

        rows = []
        for job in jobs:
            row = await FieldOptExportService.build_job_dict(job, columns, db)
            rows.append(row)

        return rows

    @staticmethod
    async def get_export_summary(
        db: AsyncSession,
        filters: dict = None,
    ) -> dict:
        """Retourne un résumé des données à exporter."""
        jobs = await FieldOptExportService.get_filtered_jobs(db, filters)

        total = len(jobs)
        completed = sum(1 for j in jobs if j.status == JobStatus.COMPLETED)
        in_progress = sum(1 for j in jobs if j.status in (
            JobStatus.IN_PROGRESS, JobStatus.EN_ROUTE
        ))
        postponed = sum(1 for j in jobs if j.status == JobStatus.POSTPONED)
        failed = sum(1 for j in jobs if j.status == JobStatus.FAILED)
        pending = sum(1 for j in jobs if j.status == JobStatus.PENDING)

        return {
            "total": total,
            "completed": completed,
            "in_progress": in_progress,
            "postponed": postponed,
            "failed": failed,
            "pending": pending,
        }

    # ── EXPORT EXCEL ──────────────────────────────────────────

    @staticmethod
    async def export_excel(
        db: AsyncSession,
        columns: List[str],
        filters: dict = None,
        include_photos: bool = False,
        include_signatures: bool = False,
    ) -> bytes:
        """Génère un fichier Excel avec les données demandées."""
        rows = await FieldOptExportService.get_jobs_data(db, filters, columns)

        wb = Workbook()
        ws = wb.active
        ws.title = "Export FieldOpt"

        # En-têtes
        if not rows:
            headers = [EXPORT_COLUMNS[c]["label"] for c in columns]
            for col_idx, header in enumerate(headers, 1):
                cell = ws.cell(row=1, column=col_idx, value=header)
                cell.fill = _header_fill(BLEU_NAVY)
                cell.font = _header_font()
                cell.alignment = _center_align()
                cell.border = _thin_border()
        else:
            first_row = rows[0]
            headers = list(first_row.keys())
            for col_idx, header in enumerate(headers, 1):
                cell = ws.cell(row=1, column=col_idx, value=header)
                cell.fill = _header_fill(BLEU_NAVY)
                cell.font = _header_font()
                cell.alignment = _center_align()
                cell.border = _thin_border()

            for row_idx, row in enumerate(rows, 2):
                for col_idx, header in enumerate(headers, 1):
                    value = row.get(header, "")
                    cell = ws.cell(row=row_idx, column=col_idx, value=value)
                    cell.border = _thin_border()
                    cell.alignment = _left_align()
                    if row_idx % 2 == 0:
                        cell.fill = _header_fill(GRIS_CLAIR)

        _autofit_columns(ws)

        ws_summary = wb.create_sheet("Résumé")
        summary = await FieldOptExportService.get_export_summary(db, filters)
        ws_summary.cell(row=1, column=1, value="Résumé de l'export").font = Font(bold=True, size=14, color=BLEU_NAVY)
        ws_summary.cell(row=3, column=1, value="Indicateur").font = Font(bold=True, color=BLEU_NAVY)
        ws_summary.cell(row=3, column=2, value="Valeur").font = Font(bold=True, color=BLEU_NAVY)
        summary_items = [
            ("Total interventions", summary["total"]),
            ("Terminées", summary["completed"]),
            ("En cours", summary["in_progress"]),
            ("Reportées", summary["postponed"]),
            ("Échecs", summary["failed"]),
            ("En attente", summary["pending"]),
        ]
        for i, (label, value) in enumerate(summary_items, 4):
            ws_summary.cell(row=i, column=1, value=label)
            ws_summary.cell(row=i, column=2, value=value)
        _autofit_columns(ws_summary)

        output = io.BytesIO()
        wb.save(output)
        return output.getvalue()

    # ── EXPORT CSV ────────────────────────────────────────────

    @staticmethod
    async def export_csv(
        db: AsyncSession,
        columns: List[str],
        filters: dict = None,
    ) -> bytes:
        """Génère un fichier CSV UTF-8 avec les données demandées."""
        rows = await FieldOptExportService.get_jobs_data(db, filters, columns)

        output = io.StringIO()
        if rows:
            writer = csv.writer(output, delimiter=";")
            writer.writerow(list(rows[0].keys()))
            for row in rows:
                writer.writerow(list(row.values()))
        else:
            headers = [EXPORT_COLUMNS[c]["label"] for c in columns]
            writer = csv.writer(output, delimiter=";")
            writer.writerow(headers)

        return output.getvalue().encode("utf-8-sig")

    # ── EXPORT PDF ────────────────────────────────────────────

    @staticmethod
    async def export_pdf(
        db: AsyncSession,
        columns: List[str],
        filters: dict = None,
        include_photos: bool = False,
        include_signatures: bool = False,
    ) -> bytes:
        """
        Génère un rapport PDF professionnel.
        Utilise une approche simple avec création de HTML puis conversion.
        """
        rows = await FieldOptExportService.get_jobs_data(db, filters, columns)

        html_parts = []

        html_parts.append("""
        <!DOCTYPE html>
        <html>
        <head>
        <meta charset="UTF-8">
        <style>
            body { font-family: 'Calibri', sans-serif; margin: 40px; color: #333; }
            h1 { color: #1F497D; font-size: 22px; border-bottom: 2px solid #1F497D; padding-bottom: 10px; }
            .header { margin-bottom: 30px; }
            .header .date { color: #666; font-size: 12px; }
            table { width: 100%; border-collapse: collapse; font-size: 11px; }
            th { background-color: #1F497D; color: white; padding: 8px; text-align: left; }
            td { padding: 6px; border-bottom: 1px solid #ddd; }
            tr:nth-child(even) { background-color: #f5f5f5; }
            .summary { margin: 20px 0; padding: 15px; background: #f0f4f8; border-radius: 5px; }
            .summary h2 { color: #1F497D; font-size: 16px; margin-top: 0; }
            .summary-grid { display: grid; grid-template-columns: repeat(3, 1fr); gap: 10px; }
            .summary-item { background: white; padding: 10px; border-radius: 4px; text-align: center; }
            .summary-item .value { font-size: 24px; font-weight: bold; color: #1F497D; }
            .summary-item .label { font-size: 11px; color: #666; }
            .footer { margin-top: 30px; font-size: 10px; color: #999; text-align: center; }
        </style>
        </head>
        <body>
        """)

        settings = get_settings()
        html_parts.append(f"""
        <div class="header">
            <h1>Rapport FieldOpt - Export FTTH</h1>
            <p class="date">Généré le {datetime.now().strftime('%d/%m/%Y à %H:%M')}</p>
            <p class="date">{settings.APP_NAME} v{settings.APP_VERSION}</p>
        </div>
        """)

        summary = await FieldOptExportService.get_export_summary(db, filters)
        html_parts.append(f"""
        <div class="summary">
            <h2>Résumé de l'export</h2>
            <div class="summary-grid">
                <div class="summary-item">
                    <div class="value">{summary["total"]}</div>
                    <div class="label">Total interventions</div>
                </div>
                <div class="summary-item">
                    <div class="value">{summary["completed"]}</div>
                    <div class="label">Terminées</div>
                </div>
                <div class="summary-item">
                    <div class="value">{summary["in_progress"]}</div>
                    <div class="label">En cours</div>
                </div>
                <div class="summary-item">
                    <div class="value">{summary["postponed"]}</div>
                    <div class="label">Reportées</div>
                </div>
                <div class="summary-item">
                    <div class="value">{summary["failed"]}</div>
                    <div class="label">Échecs</div>
                </div>
                <div class="summary-item">
                    <div class="value">{summary["pending"]}</div>
                    <div class="label">En attente</div>
                </div>
            </div>
        </div>
        """)

        if rows:
            html_parts.append("<table>")
            html_parts.append("<thead><tr>")
            for header in rows[0].keys():
                html_parts.append(f"<th>{header}</th>")
            html_parts.append("</tr></thead><tbody>")

            for row in rows:
                html_parts.append("<tr>")
                for value in row.values():
                    html_parts.append(f"<td>{value}</td>")
                html_parts.append("</tr>")

            html_parts.append("</tbody></table>")
        else:
            html_parts.append("<p>Aucune intervention trouvée pour cet export.</p>")

        html_parts.append("""
        <div class="footer">
            <p>Ce rapport est généré automatiquement par FieldOpt.</p>
        </div>
        </body></html>
        """)

        html_content = "".join(html_parts)

        try:
            import weasyprint
            pdf_bytes = weasyprint.HTML(string=html_content).write_pdf()
            return pdf_bytes
        except ImportError:
            logger.warning("weasyprint non installé, génération PDF via Fallback HTML")
            return html_content.encode("utf-8")
        except Exception as e:
            logger.error("Erreur génération PDF: %s", e)
            return html_content.encode("utf-8")

    # ── EXPORT AVEC PHOTOS (ZIP) ─────────────────────────────────

    @staticmethod
    async def export_with_photos(
        db: AsyncSession,
        columns: List[str],
        filters: dict = None,
    ) -> bytes:
        """Génère un ZIP contenant l'export Excel + les photos des interventions."""
        excel_bytes = await FieldOptExportService.export_excel(db, columns, filters)

        zip_buffer = io.BytesIO()
        with ZipFile(zip_buffer, "w") as zf:
            zf.writestr("export_fieldopt.xlsx", excel_bytes)

            jobs = await FieldOptExportService.get_filtered_jobs(db, filters)

            for job in jobs:
                if job.before_photo:
                    photo_path = job.before_photo
                    if os.path.exists(photo_path):
                        zf.write(photo_path, f"photos/{job.job_number or job.id}_before{os.path.splitext(photo_path)[1]}")
                if job.after_photo:
                    photo_path = job.after_photo
                    if os.path.exists(photo_path):
                        zf.write(photo_path, f"photos/{job.job_number or job.id}_after{os.path.splitext(photo_path)[1]}")

        return zip_buffer.getvalue()

    # ── GESTION DES MODÈLES ─────────────────────────────────────

    @staticmethod
    async def get_templates(db: AsyncSession, user_id: int = None) -> List[dict]:
        """Récupère la liste des modèles d'export."""
        query = select(ExportTemplate)
        if user_id:
            query = query.where(
                or_(ExportTemplate.created_by == user_id, ExportTemplate.is_default == True)
            )
        query = query.order_by(ExportTemplate.is_default.desc(), ExportTemplate.name)
        result = await db.execute(query)
        templates = result.scalars().all()

        return [
            {
                "id": t.id,
                "name": t.name,
                "description": t.description,
                "export_type": t.export_type,
                "columns": t.columns,
                "filters": t.filters,
                "include_photos": t.include_photos,
                "include_signatures": t.include_signatures,
                "is_default": t.is_default,
                "created_by": t.created_by,
                "created_at": t.created_at.isoformat() if t.created_at else None,
            }
            for t in templates
        ]

    @staticmethod
    async def create_template(
        db: AsyncSession,
        name: str,
        description: str = None,
        export_type: str = "excel",
        columns: List[str] = None,
        filters: dict = None,
        include_photos: bool = False,
        include_signatures: bool = True,
        is_default: bool = False,
        created_by: int = None,
    ) -> dict:
        """Crée un nouveau modèle d'export."""
        if columns is None:
            columns = [k for k, v in EXPORT_COLUMNS.items() if v["default"]]

        template = ExportTemplate(
            name=name,
            description=description,
            export_type=export_type,
            columns=columns,
            filters=filters or {},
            include_photos=include_photos,
            include_signatures=include_signatures,
            is_default=is_default,
            created_by=created_by,
        )
        db.add(template)
        await db.commit()
        await db.refresh(template)

        return {
            "id": template.id,
            "name": template.name,
            "description": template.description,
            "export_type": template.export_type,
            "columns": template.columns,
            "filters": template.filters,
            "include_photos": template.include_photos,
            "include_signatures": template.include_signatures,
            "is_default": template.is_default,
            "created_by": template.created_by,
        }

    @staticmethod
    async def update_template(
        db: AsyncSession,
        template_id: int,
        data: dict,
    ) -> dict:
        """Met à jour un modèle d'export."""
        result = await db.execute(select(ExportTemplate).where(ExportTemplate.id == template_id))
        template = result.scalar_one_or_none()
        if not template:
            raise ValueError("Modèle d'export introuvable.")

        for key in ("name", "description", "export_type", "columns", "filters",
                     "include_photos", "include_signatures", "is_default"):
            if key in data:
                setattr(template, key, data[key])

        template.updated_at = datetime.utcnow()
        await db.commit()
        await db.refresh(template)

        return {
            "id": template.id,
            "name": template.name,
            "description": template.description,
            "export_type": template.export_type,
            "columns": template.columns,
            "filters": template.filters,
            "include_photos": template.include_photos,
            "include_signatures": template.include_signatures,
            "is_default": template.is_default,
        }

    @staticmethod
    async def delete_template(db: AsyncSession, template_id: int):
        """Supprime un modèle d'export."""
        result = await db.execute(select(ExportTemplate).where(ExportTemplate.id == template_id))
        template = result.scalar_one_or_none()
        if not template:
            raise ValueError("Modèle d'export introuvable.")
        await db.delete(template)
        await db.commit()

    # ── HISTORIQUE DES EXPORTS ──────────────────────────────────

    @staticmethod
    async def log_export(
        db: AsyncSession,
        user_id: int = None,
        template_id: int = None,
        export_name: str = None,
        export_format: str = "excel",
        job_count: int = 0,
        filters_used: dict = None,
        columns_used: list = None,
        has_photos: bool = False,
        has_signatures: bool = False,
        file_size_bytes: int = None,
        file_path: str = None,
        status: str = "completed",
        error_message: str = None,
        duration_seconds: float = None,
    ):
        """Enregistre un export dans l'historique."""
        history = ExportHistory(
            user_id=user_id,
            template_id=template_id,
            export_name=export_name or f"Export {datetime.now().strftime('%Y%m%d_%H%M%S')}",
            export_format=export_format,
            job_count=job_count,
            filters_used=filters_used or {},
            columns_used=columns_used or [],
            has_photos=has_photos,
            has_signatures=has_signatures,
            file_size_bytes=file_size_bytes,
            file_path=file_path,
            status=status,
            error_message=error_message,
            duration_seconds=duration_seconds,
        )
        db.add(history)
        await db.commit()

    @staticmethod
    async def get_export_history(
        db: AsyncSession,
        limit: int = 50,
        offset: int = 0,
        user_id: int = None,
    ) -> List[dict]:
        """Récupère l'historique des exports."""
        query = select(ExportHistory).order_by(ExportHistory.created_at.desc())
        if user_id:
            query = query.where(ExportHistory.user_id == user_id)
        query = query.offset(offset).limit(limit)
        result = await db.execute(query)
        records = result.scalars().all()

        return [
            {
                "id": r.id,
                "user_id": r.user_id,
                "template_id": r.template_id,
                "export_name": r.export_name,
                "export_format": r.export_format,
                "job_count": r.job_count,
                "filters_used": r.filters_used,
                "columns_used": r.columns_used,
                "has_photos": r.has_photos,
                "has_signatures": r.has_signatures,
                "file_size_bytes": r.file_size_bytes,
                "status": r.status,
                "error_message": r.error_message,
                "duration_seconds": r.duration_seconds,
                "created_at": r.created_at.isoformat() if r.created_at else None,
            }
            for r in records
        ]

    # ── PROFILS PRÉDÉFINIS ───────────────────────────────────────

    @staticmethod
    def get_available_profiles() -> dict:
        """Retourne les profils d'export prédéfinis."""
        return EXPORT_PROFILES

    @staticmethod
    def get_available_columns() -> dict:
        """Retourne les colonnes disponibles pour l'export."""
        return EXPORT_COLUMNS

    @staticmethod
    def get_column_categories() -> dict:
        """Retourne les catégories de colonnes."""
        categories = {}
        for key, col in EXPORT_COLUMNS.items():
            cat = col.get("category", "autres")
            if cat not in categories:
                categories[cat] = {"label": cat.capitalize(), "columns": []}
            categories[cat]["columns"].append({
                "key": key,
                "label": col["label"],
                "default": col["default"],
            })
        return categories
