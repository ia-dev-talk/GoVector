from pathlib import Path
import logging
import csv
import io

from openpyxl import load_workbook

logger = logging.getLogger(__name__)


def color_to_hex(color):

    if color is None:
        return None

    try:

        rgb = color.rgb

        if rgb is None:
            return None

        if len(rgb) == 8:
            rgb = rgb[2:]

        return "#" + rgb.upper()

    except Exception:

        return None


def parse_excel(file_path: str):
    logger.info(f"Démarrage du parsing Excel pour le fichier: {file_path}")

    wb = load_workbook(

        filename=file_path,

        data_only=True  # Lire les valeurs calculées au lieu des formules

    )

    workbook = []

    for sheet in wb.worksheets:

        merged_ranges = {
            str(rng)
            for rng in sheet.merged_cells.ranges
        }

        rows = []

        for row in sheet.iter_rows():

            current = []

            for cell in row:
                logger.debug(f"Parsing de la cellule {cell.coordinate} sur la feuille '{sheet.title}': Valeur='{cell.value}', Type='{type(cell.value)}'")
                fill = color_to_hex(cell.fill.fgColor)

                font_color = color_to_hex(cell.font.color)

                current.append({

                    "row": cell.row,

                    "column": cell.column,

                    "coordinate": cell.coordinate,

                    "value": cell.value,

                    "color": fill,

                    "font_color": font_color,

                    "bold": bool(cell.font.bold),

                    "italic": bool(cell.font.italic),

                    "font_size": cell.font.sz,

                    "horizontal_align": cell.alignment.horizontal,

                    "vertical_align": cell.alignment.vertical,

                    "number_format": cell.number_format,

                    "comment": (
                        cell.comment.text
                        if cell.comment
                        else None
                    ),

                    "merged": any(
                        cell.coordinate in rng
                        for rng in merged_ranges
                    )

                })

            rows.append(current)

        workbook.append({

            "sheet": sheet.title,

            "rows": rows,

            "max_row": sheet.max_row,

            "max_column": sheet.max_column,

            "merged_ranges": list(merged_ranges)

        })

    return workbook


def parse_csv(file_path: str):
    logger.info(f"Démarrage du parsing CSV pour le fichier: {file_path}")

    workbook = []

    raw_content = Path(file_path).read_bytes()
    try:
        text_content = raw_content.decode("utf-8-sig")
    except UnicodeDecodeError:
        # Les exports bureautiques Windows francophones utilisent encore
        # fréquemment CP-1252. Le décodage reste déterministe et ne modifie
        # jamais le contenu métier.
        text_content = raw_content.decode("cp1252")

    sample = text_content[:8192]
    try:
        dialect = csv.Sniffer().sniff(sample, delimiters=",;\t|")
    except csv.Error:
        dialect = csv.excel

    reader = csv.reader(io.StringIO(text_content, newline=""), dialect)
    rows_data = list(reader)

    if not rows_data:
        return workbook

    rows = []

    for row_index, row_values in enumerate(rows_data, start=1):

        current = []

        for col_index, value in enumerate(row_values, start=1):
            logger.debug(f"Parsing de la cellule {row_index}:{col_index} sur la feuille \'{Path(file_path).stem}\': Valeur=\'{value}\', Type=\'{type(value)}\'")
            text = value.strip() if isinstance(value, str) else value

            if text == "":
                text = None

            current.append({
                "row": row_index,
                "column": col_index,
                "coordinate": f"{row_index}:{col_index}",
                "value": text,
                "color": None,
                "font_color": None,
                "bold": False,
                "italic": False,
                "font_size": None,
                "horizontal_align": None,
                "vertical_align": None,
                "number_format": None,
                "comment": None,
                "merged": False,
            })

        rows.append(current)

    workbook.append({
        "sheet": Path(file_path).stem,
        "rows": rows,
        "max_row": len(rows),
        "max_column": max(len(r) for r in rows_data),
        "merged_ranges": [],
    })

    return workbook


def parse_spreadsheet(file_path: str):

    extension = Path(file_path).suffix.lower()

    if extension == ".csv":
        return parse_csv(file_path)

    return parse_excel(file_path)
